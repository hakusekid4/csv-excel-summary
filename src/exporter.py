"""集計結果を Excel ファイル(bytes)にする。画面のダウンロードボタンと CLI の両方から使う。"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")  # 見出し行の背景(薄い青)
MAX_SHEET_NAME = 31  # Excel のシート名は 31 文字まで


def _safe_sheet_name(name: str) -> str:
    """Excel で使えないシート名を直す(禁止文字を消し、31 文字に切る)。"""
    for ch in '[]:*?/\\':
        name = name.replace(ch, "")
    return (name or "Sheet")[:MAX_SHEET_NAME]


def to_excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    """{シート名: DataFrame} を 1 つの Excel にまとめ、bytes で返す。

    見出しを太字+色付きにし、列幅を内容に合わせて広げる。
    """
    if not sheets:
        raise ValueError("出力するシートがありません。")
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            sheet_name = _safe_sheet_name(name)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL
            for i, col in enumerate(df.columns, start=1):
                # 列幅 = その列で一番長い文字列の長さ(全角を考えて 1.8 倍)+ 余白
                longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
                ws.column_dimensions[get_column_letter(i)].width = min(60, longest * 1.8 + 2)
            ws.freeze_panes = "A2"  # 見出し行を固定
    return buf.getvalue()
